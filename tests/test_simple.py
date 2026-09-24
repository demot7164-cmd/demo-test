"""
Simple integration test for DOCX translation workflow.
Tests the core functionality without workflow orchestration.
"""

import io
import tempfile
from pathlib import Path

from docx import Document

from src.workflows.activities import (
    extract_docx_content,
    translate_content_chunks,
    rebuild_docx,
)
from src.models.schemas import (
    TranslationInput,
    Language,
    TranslationChunk,
    DOCXElementType,
    DOCXContent,
)


def test_extract_and_rebuild():
    """Test extraction and rebuilding of DOCX content."""
    # Create a DOCX with various elements
    doc = Document()
    doc.add_heading('Test Document', 0)
    doc.add_paragraph('This is a test paragraph.')
    doc.add_paragraph('This is another paragraph.')
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Header 1'
    table.cell(0, 1).text = 'Header 2'
    table.cell(1, 0).text = 'Data 1'
    table.cell(1, 1).text = 'Data 2'
    
    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    docx_bytes = buffer.read()
    
    # Create translation input
    input_data = TranslationInput(
        docx_bytes=docx_bytes,
        target_language=Language.FRENCH,
        preserve_formatting=True,
        translate_headers_footers=True,
        translate_table_text=True,
    )
    
    # Test extraction (synchronous for testing)
    import asyncio
    
    async def test_async():
        content = await extract_docx_content(input_data)
        
        # Verify extraction
        assert len(content.paragraphs) >= 2  # heading + paragraphs
        assert len(content.tables) == 1
        
        # Create simple translation chunks
        chunks = []
        for i, para in enumerate(content.paragraphs):
            if para.text.strip():
                chunks.append(TranslationChunk(
                    text=para.text,
                    element_type=DOCXElementType.PARAGRAPH,
                    element_index=i,
                ))
        
        # Test translation (mock)
        translated_chunks = await translate_content_chunks(chunks, Language.FRENCH)
        
        # Test rebuild
        result_bytes = await rebuild_docx(content, translated_chunks)
        
        # Verify result
        assert isinstance(result_bytes, bytes)
        assert len(result_bytes) > 0
        
        # Save and verify it's a valid DOCX
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            f.write(result_bytes)
            temp_path = f.name
        
        try:
            result_doc = Document(temp_path)
            assert len(result_doc.paragraphs) > 0
            
            # Check that translation is present
            full_text = '\n'.join([p.text for p in result_doc.paragraphs])
            assert '[Translated' in full_text
            
        finally:
            Path(temp_path).unlink(missing_ok=True)
    
    asyncio.run(test_async())
    print("✓ Extract and rebuild test passed!")


def test_table_preservation():
    """Test that tables are preserved correctly."""
    # Create DOCX with table
    doc = Document()
    
    # Add a 3x3 table
    table = doc.add_table(rows=3, cols=3)
    for i in range(3):
        for j in range(3):
            table.cell(i, j).text = f'Cell {i},{j}'
    
    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    docx_bytes = buffer.read()
    
    # Test extraction and rebuild
    import asyncio
    
    async def test_async():
        input_data = TranslationInput(
            docx_bytes=docx_bytes,
            target_language=Language.SPANISH,
        )
        
        content = await extract_docx_content(input_data)
        
        # Verify table extraction
        assert len(content.tables) == 1
        assert len(content.tables[0]) == 3  # 3 rows
        assert len(content.tables[0][0]) == 3  # 3 cols in first row
        
        # Create translation chunks for table
        chunks = []
        for table_idx, table_rows in enumerate(content.tables):
            for row_idx, row_cells in enumerate(table_rows):
                for col_idx, cell in enumerate(row_cells):
                    chunks.append(TranslationChunk(
                        text=cell.text,
                        element_type=DOCXElementType.TABLE,
                        element_index=table_idx,
                        sub_element_index=row_idx * 100 + col_idx,
                    ))
        
        # Translate
        translated_chunks = await translate_content_chunks(chunks, Language.SPANISH)
        
        # Rebuild
        result_bytes = await rebuild_docx(content, translated_chunks)
        
        # Verify result has table
        result_doc = Document(io.BytesIO(result_bytes))
        assert len(result_doc.tables) == 1
        assert len(result_doc.tables[0].rows) == 3
        
        # Check translation in table
        table_text = []
        for row in result_doc.tables[0].rows:
            for cell in row.cells:
                table_text.append(cell.text)
        
        full_table_text = '\n'.join(table_text)
        assert '[Translated' in full_table_text
    
    asyncio.run(test_async())
    print("✓ Table preservation test passed!")


if __name__ == '__main__':
    test_extract_and_rebuild()
    test_table_preservation()
    print("\n✅ All tests passed!")
